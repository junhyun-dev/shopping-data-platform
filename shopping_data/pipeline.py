from __future__ import annotations

import hashlib
import json
import os
import uuid
import fcntl
from collections.abc import Callable
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, TextIO

import duckdb
from openpyxl import load_workbook

from .config import (
    AGGREGATE_SQL,
    CONTRIBUTIONS_SQL,
    SCHEMA_SQL,
    TRANSFORMATION_SQL,
)


EXPECTED_HEADERS = (
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
)

UCI_DATASET_URL = "https://archive.ics.uci.edu/dataset/352/online+retail"


class PipelineFailure(RuntimeError):
    """The input could not produce a publishable result."""


class PipelineBusy(PipelineFailure):
    """Another live local writer owns the database pipeline lock."""


class ResultNotFound(LookupError):
    """No successfully published result exists yet."""


class ResultReadError(RuntimeError):
    """A published-result pointer exists but cannot be read safely."""


ProcessCheckpoint = Callable[[str, str], None]

PROCESS_CHECKPOINTS = (
    "before_data_commit",
    "after_data_commit",
    "after_artifact_write",
    "after_publish_commit",
)


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _completion_time(started_at: datetime) -> datetime:
    """Keep one run's completion from preceding its own start if the wall clock shifts."""
    return max(utc_now(), started_at)


def _writer_lock_path(database_path: Path) -> Path:
    return database_path.with_name(f"{database_path.name}.writer.lock")


def _acquire_writer_lock(
    database_path: Path, owner_token: str
) -> tuple[TextIO, Path]:
    lock_path = _writer_lock_path(database_path)
    lock_file = lock_path.open("a+", encoding="utf-8")
    try:
        fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError as exc:
        lock_file.close()
        raise PipelineBusy(
            f"another local pipeline writer holds {lock_path}"
        ) from exc

    lock_file.seek(0)
    lock_file.truncate()
    json.dump(
        {
            "owner_pid": os.getpid(),
            "owner_token": owner_token,
            "acquired_at": utc_now().isoformat(timespec="seconds"),
        },
        lock_file,
    )
    lock_file.write("\n")
    lock_file.flush()
    os.fsync(lock_file.fileno())
    return lock_file, lock_path


def _release_writer_lock(lock_file: TextIO) -> None:
    try:
        fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)
    finally:
        lock_file.close()


def _invoke_process_checkpoint(
    process_checkpoint: ProcessCheckpoint | None,
    name: str,
    run_id: str,
) -> None:
    if process_checkpoint is not None:
        process_checkpoint(name, run_id)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_manifest(path: Path) -> dict[str, Any]:
    manifest = json.loads(path.read_text(encoding="utf-8"))
    required = {
        "dataset",
        "source_url",
        "expected_sha256",
        "sheet",
        "selected_excel_rows",
        "scope_note",
    }
    missing = sorted(required - manifest.keys())
    if missing:
        raise PipelineFailure(f"source manifest is missing: {', '.join(missing)}")

    manifest["selected_excel_rows"] = _normalize_selected_rows(
        manifest["selected_excel_rows"]
    )
    return manifest


def _normalize_selected_rows(selected_rows: list[int]) -> list[int]:
    if not selected_rows or any(
        not isinstance(row, int) or isinstance(row, bool) or row < 2
        for row in selected_rows
    ):
        raise PipelineFailure("selected_excel_rows must contain Excel data-row numbers")
    if len(selected_rows) != len(set(selected_rows)):
        raise PipelineFailure("selected_excel_rows contains duplicate row numbers")
    return sorted(selected_rows)


def _stable_id(payload: Any) -> str:
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return f"sha256:{hashlib.sha256(encoded).hexdigest()}"


def _input_scope_id(
    source_file_id: str, sheet_name: str, selected_rows: list[int]
) -> str:
    return _stable_id(
        {
            "source_file_id": source_file_id,
            "sheet_name": sheet_name,
            "selected_excel_rows": selected_rows,
        }
    )


def _sql_rule_id(sql_text: str) -> str:
    return f"sha256:{hashlib.sha256(sql_text.encode('utf-8')).hexdigest()}"


def _raw_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value)


def _invoice_no_source_kind(value: Any) -> str:
    if value is None:
        return "missing"
    if isinstance(value, bool):
        return "unsupported"
    if isinstance(value, str):
        return "text"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float) and value.is_integer():
        return "integer"
    return "unsupported"


def read_selected_rows(
    workbook_path: Path, sheet_name: str, selected_rows: list[int]
) -> list[dict[str, str | int | None]]:
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl exposes several format-specific exception types
        raise PipelineFailure(f"workbook could not be read: {exc}") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise PipelineFailure(f"workbook sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        wanted = set(selected_rows)
        found: dict[int, dict[str, str | int | None]] = {}
        header: tuple[Any, ...] | None = None

        for row_number, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=max(selected_rows), max_col=8, values_only=True),
            start=1,
        ):
            if row_number == 1:
                header = tuple(values)
                if header != EXPECTED_HEADERS:
                    raise PipelineFailure(
                        f"unexpected workbook header: expected {EXPECTED_HEADERS}, got {header}"
                    )
                continue
            if row_number not in wanted:
                continue
            found[row_number] = {
                "source_row_number": row_number,
                "invoice_no_raw": _raw_text(values[0]),
                "invoice_no_source_kind": _invoice_no_source_kind(values[0]),
                "stock_code_raw": _raw_text(values[1]),
                "description_raw": _raw_text(values[2]),
                "quantity_raw": _raw_text(values[3]),
                "invoice_date_raw": _raw_text(values[4]),
                "unit_price_raw": _raw_text(values[5]),
                "customer_id_raw": _raw_text(values[6]),
                "country_raw": _raw_text(values[7]),
            }

        missing = sorted(wanted - found.keys())
        if missing:
            raise PipelineFailure(f"selected Excel rows were not found: {missing}")
        return [found[row_number] for row_number in selected_rows]
    finally:
        workbook.close()


def initialize_database(connection: duckdb.DuckDBPyConnection) -> None:
    connection.execute(SCHEMA_SQL.read_text(encoding="utf-8"))
    connection.execute(TRANSFORMATION_SQL.read_text(encoding="utf-8"))


def _preserve_unpublished_artifact(
    published_root: Path, run_id: str
) -> Path | None:
    artifact = (published_root / f"{run_id}.json").resolve()
    orphaned = (published_root / "orphaned" / f"{run_id}.json").resolve()
    if not artifact.is_relative_to(published_root) or not orphaned.is_relative_to(
        published_root
    ):
        raise PipelineFailure("recovery artifact path escaped the published directory")
    if artifact.exists() and orphaned.exists():
        raise PipelineFailure(
            f"both active and orphaned artifacts exist for interrupted run {run_id}"
        )
    if artifact.exists():
        if not artifact.is_file():
            raise PipelineFailure(
                f"interrupted run artifact is not a file: {artifact}"
            )
        orphaned.parent.mkdir(parents=True, exist_ok=True)
        os.replace(artifact, orphaned)
        return orphaned
    if orphaned.is_file():
        return orphaned
    return None


def _recover_orphaned_runs(
    connection: duckdb.DuckDBPyConnection,
    published_root: Path,
) -> list[str]:
    """Recover only after the caller owns the exclusive local writer lock.

    The lock, not elapsed time, proves that no live cooperating writer still owns
    these ``running`` rows. Unpublished artifacts are preserved under their exact
    run IDs and remain outside ``published_results``.
    """

    rows = connection.execute(
        """
        SELECT r.run_id, r.started_at, COUNT(pr.run_id) AS published_count
        FROM pipeline_runs AS r
        LEFT JOIN published_results AS pr USING (run_id)
        WHERE r.status = 'running'
        GROUP BY r.run_id, r.started_at
        ORDER BY r.run_id
        """
    ).fetchall()
    if not rows:
        return []

    recovered: list[tuple[str, datetime, str, str | None]] = []
    for run_id, started_at, published_count in rows:
        if published_count:
            raise PipelineFailure(
                f"running run already has published evidence and needs manual review: {run_id}"
            )
        preserved_artifact = _preserve_unpublished_artifact(published_root, run_id)
        reason = (
            "previous writer process ended before publication completed; "
            "recovered by the next exclusive local writer"
        )
        recovered.append(
            (
                run_id,
                _completion_time(started_at),
                reason,
                str(preserved_artifact) if preserved_artifact else None,
            )
        )

    connection.execute("BEGIN TRANSACTION")
    try:
        connection.executemany(
            """
            UPDATE pipeline_runs
            SET status = 'failed', completed_at = ?, error_message = ?,
                recovered_at = ?, recovery_reason = ?,
                recovered_artifact_path = ?
            WHERE run_id = ? AND status = 'running'
            """,
            [
                [completed_at, reason, utc_now(), reason, artifact_path, run_id]
                for run_id, completed_at, reason, artifact_path in recovered
            ],
        )
        connection.execute("COMMIT")
    except Exception:
        connection.execute("ROLLBACK")
        raise
    return [row[0] for row in recovered]


def _validate_typed_rows(
    connection: duckdb.DuckDBPyConnection,
    run_id: str,
    expected_count: int,
) -> None:
    actual_count = connection.execute(
        "SELECT COUNT(*) FROM run_input_rows WHERE run_id = ?", [run_id]
    ).fetchone()[0]
    if actual_count != expected_count:
        raise PipelineFailure(
            f"source row count mismatch: expected {expected_count}, got {actual_count}"
        )

    quantity_rows = connection.execute(
        """
        SELECT t.source_row_number, t.quantity_source_text, t.quantity
        FROM typed_transaction_lines AS t
        JOIN run_input_rows AS i
          ON i.source_file_id = t.source_file_id
         AND i.sheet_name = t.sheet_name
         AND i.source_row_number = t.source_row_number
        WHERE i.run_id = ?
        ORDER BY t.source_row_number
        """,
        [run_id],
    ).fetchall()
    lossy_quantity_rows: list[int] = []
    for source_row_number, quantity_raw, quantity in quantity_rows:
        try:
            source_quantity = Decimal(quantity_raw)
            is_lossless_integer = (
                source_quantity.is_finite()
                and source_quantity == source_quantity.to_integral_value()
                and quantity is not None
                and source_quantity == Decimal(quantity)
            )
        except (InvalidOperation, TypeError, ValueError):
            is_lossless_integer = False
        if not is_lossless_integer:
            lossy_quantity_rows.append(int(source_row_number))

    if lossy_quantity_rows:
        raise PipelineFailure(
            "quantity cannot be represented losslessly as BIGINT at Excel rows: "
            + ", ".join(str(row) for row in lossy_quantity_rows)
        )

    invalid = connection.execute(
        """
        SELECT t.source_row_number
        FROM typed_transaction_lines AS t
        JOIN run_input_rows AS i
          ON i.source_file_id = t.source_file_id
         AND i.sheet_name = t.sheet_name
         AND i.source_row_number = t.source_row_number
        WHERE i.run_id = ?
          AND (
            t.stock_code IS NULL
            OR t.invoice_timestamp IS NULL
            OR t.unit_price_source_text IS NULL
          )
        ORDER BY t.source_row_number
        """,
        [run_id],
    ).fetchall()
    if invalid:
        raise PipelineFailure(
            "typed fields are invalid at Excel rows: "
            + ", ".join(str(row[0]) for row in invalid)
        )


def _format_timestamp(value: datetime) -> str:
    return value.isoformat(sep=" ", timespec="seconds")


def _build_result(
    connection: duckdb.DuckDBPyConnection,
    *,
    run_id: str,
    source_file_id: str,
    source_sha256: str,
    source_path: Path,
    manifest: dict[str, Any],
    selected_rows: list[int],
    input_scope_id: str,
    aggregation_rule_id: str,
    transformation_rule_id: str,
    completed_at: datetime,
) -> dict[str, Any]:
    aggregate_sql = AGGREGATE_SQL.read_text(encoding="utf-8").strip()
    transformation_sql = TRANSFORMATION_SQL.read_text(encoding="utf-8").strip()
    contribution_sql = CONTRIBUTIONS_SQL.read_text(encoding="utf-8").strip()
    aggregate_rows = connection.execute(aggregate_sql, [run_id]).fetchall()

    products: list[dict[str, Any]] = []
    for _, observed_date, stock_code, quantity_sum, row_count in aggregate_rows:
        contributions = connection.execute(
            contribution_sql, [run_id, stock_code, observed_date]
        ).fetchall()
        products.append(
            {
                "data_state": "available",
                "observed_date": observed_date.isoformat(),
                "stock_code": stock_code,
                "sample_quantity_sum": int(quantity_sum),
                "observed_row_count": int(row_count),
                "contributing_rows": [
                    {
                        "sheet": row[0],
                        "source_row_number": int(row[1]),
                        "invoice_no": row[2],
                        "source_cancellation_marker": row[3],
                        "invoice_no_source_kind": row[4],
                        "stock_code": row[5],
                        "description": row[6],
                        "quantity": int(row[7]),
                        "invoice_timestamp": _format_timestamp(row[8]),
                        "unit_price_source_text": row[9],
                        "country": row[10],
                    }
                    for row in contributions
                ],
            }
        )

    return {
        "result_state": "complete_for_selected_rows",
        "run": {
            "run_id": run_id,
            "status": "succeeded",
            "completed_at": completed_at.isoformat(timespec="seconds"),
        },
        "source": {
            "dataset": manifest["dataset"],
            "source_url": manifest["source_url"],
            "source_file_id": source_file_id,
            "sha256": source_sha256,
            "stored_file": str(source_path.resolve()),
            "sheet": manifest["sheet"],
            "selected_excel_rows": selected_rows,
            "input_scope_id": input_scope_id,
            "scope_note": manifest["scope_note"],
        },
        "metric_contract": {
            "sample_quantity_sum": "Signed Quantity summed only across the selected source rows.",
            "observed_row_count": "Number of selected source rows in the product/date group.",
            "date_basis": "Date portion of the source-displayed InvoiceDate; not a business-day or timezone contract.",
            "not_claimed": [
                "whole-day or full-dataset result",
                "sales, cancellation, refund, payment, or net-revenue metric",
                "business duplicate detection",
            ],
        },
        "aggregation_rule_id": aggregation_rule_id,
        "transformation_rule_id": transformation_rule_id,
        "source_cancellation_marker_contract": {
            "source_field": "InvoiceNo",
            "rule": (
                "Preserve InvoiceNo text. For this derived observation only, remove "
                "surrounding ASCII space, tab, CR, and LF, then compare the first "
                "character case-insensitively with C. Missing, supported-whitespace-only, "
                "or unsupported source-cell types are unknown."
            ),
            "states": {
                "true": "The observed InvoiceNo starts with C/c.",
                "false": "No C/c prefix was observed; this does not confirm a sale.",
                "unknown": (
                    "InvoiceNo was missing, contained only ASCII space/tab/CR/LF, "
                    "or had an unsupported source-cell type."
                ),
            },
            "official_definition_url": UCI_DATASET_URL,
            "not_claimed": [
                "original transaction link",
                "refund or payment state",
                "discount classification",
            ],
        },
        "applied_transformation_sql": transformation_sql,
        "applied_sql": aggregate_sql,
        "products": products,
    }


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with temporary.open("x", encoding="utf-8") as output:
            json.dump(payload, output, ensure_ascii=False, indent=2)
            output.write("\n")
            output.flush()
            os.fsync(output.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def run_pipeline(
    *,
    source_path: Path,
    manifest_path: Path,
    database_path: Path,
    published_root: Path,
    selected_rows: list[int] | None = None,
    failure_point: str | None = None,
    process_checkpoint: ProcessCheckpoint | None = None,
) -> dict[str, Any]:
    """Run one idempotent sample load and atomically publish a successful result.

    ``failure_point`` is an internal test seam. It is intentionally not exposed by the CLI.
    """

    allowed_failure_points = {None, "after_data", "after_artifact"}
    if failure_point not in allowed_failure_points:
        raise ValueError(f"unknown failure point: {failure_point}")

    manifest = load_manifest(manifest_path)
    selected_rows = _normalize_selected_rows(
        selected_rows if selected_rows is not None else manifest["selected_excel_rows"]
    )
    source_path = source_path.resolve()
    database_path = database_path.resolve()
    published_root = published_root.resolve()
    database_path.parent.mkdir(parents=True, exist_ok=True)
    published_root.mkdir(parents=True, exist_ok=True)

    source_sha256 = sha256_file(source_path)
    source_file_id = f"sha256:{source_sha256}"
    input_scope_id = _input_scope_id(
        source_file_id, manifest["sheet"], selected_rows
    )
    aggregate_sql = AGGREGATE_SQL.read_text(encoding="utf-8").strip()
    aggregation_rule_id = _sql_rule_id(aggregate_sql)
    transformation_sql = TRANSFORMATION_SQL.read_text(encoding="utf-8").strip()
    transformation_rule_id = _sql_rule_id(transformation_sql)
    run_id = str(uuid.uuid4())
    owner_token = f"{os.getpid()}:{uuid.uuid4()}"
    started_at = utc_now()
    artifact_path: Path | None = None
    published = False

    lock_file, _ = _acquire_writer_lock(database_path, owner_token)
    connection: duckdb.DuckDBPyConnection | None = None
    try:
        connection = duckdb.connect(str(database_path))
        initialize_database(connection)
        _recover_orphaned_runs(connection, published_root)
        connection.execute(
            """
            INSERT INTO pipeline_runs
                (run_id, source_file_id, input_path, status, started_at,
                 owner_pid, owner_token)
            VALUES (?, ?, ?, 'running', ?, ?, ?)
            """,
            [
                run_id,
                source_file_id,
                str(source_path),
                started_at,
                os.getpid(),
                owner_token,
            ],
        )

        try:
            expected_sha256 = manifest["expected_sha256"]
            if source_sha256 != expected_sha256:
                raise PipelineFailure(
                    f"source SHA-256 mismatch: expected {expected_sha256}, got {source_sha256}"
                )

            connection.execute(
                """
                INSERT INTO run_input_scopes
                    (run_id, source_file_id, sheet_name, selected_excel_rows,
                     input_scope_id, aggregation_rule_id, transformation_rule_id)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    run_id,
                    source_file_id,
                    manifest["sheet"],
                    json.dumps(selected_rows, separators=(",", ":")),
                    input_scope_id,
                    aggregation_rule_id,
                    transformation_rule_id,
                ],
            )
            rows = read_selected_rows(source_path, manifest["sheet"], selected_rows)
            loaded_at = utc_now()

            connection.execute("BEGIN TRANSACTION")
            connection.execute(
                """
                INSERT OR IGNORE INTO source_files
                    (source_file_id, sha256, original_file_name, source_url,
                     stored_path, workbook_sheet, selected_excel_rows, registered_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    source_file_id,
                    source_sha256,
                    source_path.name,
                    manifest["source_url"],
                    str(source_path),
                    manifest["sheet"],
                    json.dumps(selected_rows, separators=(",", ":")),
                    loaded_at,
                ],
            )
            connection.executemany(
                """
                INSERT OR IGNORE INTO source_rows
                    (source_file_id, sheet_name, source_row_number,
                     invoice_no_raw, invoice_no_source_kind,
                     stock_code_raw, description_raw,
                     quantity_raw, invoice_date_raw, unit_price_raw,
                     customer_id_raw, country_raw, loaded_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    [
                        source_file_id,
                        manifest["sheet"],
                        row["source_row_number"],
                        row["invoice_no_raw"],
                        row["invoice_no_source_kind"],
                        row["stock_code_raw"],
                        row["description_raw"],
                        row["quantity_raw"],
                        row["invoice_date_raw"],
                        row["unit_price_raw"],
                        row["customer_id_raw"],
                        row["country_raw"],
                        loaded_at,
                    ]
                    for row in rows
                ],
            )
            connection.executemany(
                """
                UPDATE source_rows
                SET invoice_no_source_kind = COALESCE(invoice_no_source_kind, ?)
                WHERE source_file_id = ?
                  AND sheet_name = ?
                  AND source_row_number = ?
                """,
                [
                    [
                        row["invoice_no_source_kind"],
                        source_file_id,
                        manifest["sheet"],
                        row["source_row_number"],
                    ]
                    for row in rows
                ],
            )
            connection.executemany(
                """
                INSERT INTO run_input_rows
                    (run_id, source_file_id, sheet_name, source_row_number)
                VALUES (?, ?, ?, ?)
                """,
                [
                    [
                        run_id,
                        source_file_id,
                        manifest["sheet"],
                        row_number,
                    ]
                    for row_number in selected_rows
                ],
            )
            _validate_typed_rows(connection, run_id, len(selected_rows))
            _invoke_process_checkpoint(
                process_checkpoint, "before_data_commit", run_id
            )
            connection.execute("COMMIT")
            _invoke_process_checkpoint(
                process_checkpoint, "after_data_commit", run_id
            )

            if failure_point == "after_data":
                raise PipelineFailure("simulated failure after database load")

            completed_at = _completion_time(started_at)
            payload = _build_result(
                connection,
                run_id=run_id,
                source_file_id=source_file_id,
                source_sha256=source_sha256,
                source_path=source_path,
                manifest=manifest,
                selected_rows=selected_rows,
                input_scope_id=input_scope_id,
                aggregation_rule_id=aggregation_rule_id,
                transformation_rule_id=transformation_rule_id,
                completed_at=completed_at,
            )
            artifact_path = published_root / f"{run_id}.json"
            _write_json_atomic(artifact_path, payload)
            _invoke_process_checkpoint(
                process_checkpoint, "after_artifact_write", run_id
            )

            if failure_point == "after_artifact":
                raise PipelineFailure("simulated failure after result artifact write")

            connection.execute("BEGIN TRANSACTION")
            connection.execute(
                "UPDATE published_results SET is_current = FALSE WHERE is_current = TRUE"
            )
            connection.execute(
                """
                INSERT INTO published_results
                    (run_id, artifact_path, published_at, is_current)
                VALUES (?, ?, ?, TRUE)
                """,
                [run_id, str(artifact_path), completed_at],
            )
            connection.execute(
                """
                UPDATE pipeline_runs
                SET status = 'succeeded', completed_at = ?, error_message = NULL
                WHERE run_id = ?
                """,
                [completed_at, run_id],
            )
            connection.execute("COMMIT")
            published = True
            _invoke_process_checkpoint(
                process_checkpoint, "after_publish_commit", run_id
            )
            return payload
        except Exception as exc:
            if published:
                raise
            try:
                connection.execute("ROLLBACK")
            except duckdb.TransactionException:
                pass
            if artifact_path is not None and not published:
                artifact_path.unlink(missing_ok=True)
            connection.execute(
                """
                UPDATE pipeline_runs
                SET status = 'failed', completed_at = ?, error_message = ?
                WHERE run_id = ?
                """,
                [_completion_time(started_at), str(exc), run_id],
            )
            raise
    finally:
        if connection is not None:
            connection.close()
        _release_writer_lock(lock_file)


def load_current_result(database_path: Path, published_root: Path) -> dict[str, Any]:
    database_path = database_path.resolve()
    published_root = published_root.resolve()
    if not database_path.exists():
        raise ResultNotFound("no local result database exists")

    try:
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            current = connection.execute(
                """
                SELECT pr.artifact_path
                FROM published_results pr
                JOIN pipeline_runs r USING (run_id)
                WHERE pr.is_current = TRUE AND r.status = 'succeeded'
                """
            ).fetchall()
        finally:
            connection.close()
    except Exception as exc:
        raise ResultReadError(f"result database could not be read: {exc}") from exc

    if not current:
        raise ResultNotFound("no successful result has been published")
    if len(current) != 1:
        raise ResultReadError("more than one current result is registered")

    artifact = Path(current[0][0]).resolve()
    if not artifact.is_relative_to(published_root):
        raise ResultReadError("published artifact points outside the result directory")
    try:
        return json.loads(artifact.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ResultReadError(f"published result could not be read: {exc}") from exc


def _load_published_result(
    database_path: Path,
    published_root: Path,
    run_id: str,
) -> tuple[dict[str, Any], bool]:
    database_path = database_path.resolve()
    published_root = published_root.resolve()
    if not database_path.exists():
        raise ResultNotFound("no local result database exists")

    try:
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            row = connection.execute(
                """
                SELECT pr.artifact_path, pr.is_current
                FROM published_results AS pr
                JOIN pipeline_runs AS r USING (run_id)
                WHERE pr.run_id = ? AND r.status = 'succeeded'
                """,
                [run_id],
            ).fetchone()
        finally:
            connection.close()
    except Exception as exc:
        raise ResultReadError(f"result database could not be read: {exc}") from exc

    if row is None:
        raise ResultNotFound(f"successful published result not found: {run_id}")
    artifact = Path(row[0]).resolve()
    if not artifact.is_relative_to(published_root):
        raise ResultReadError("published artifact points outside the result directory")
    try:
        return json.loads(artifact.read_text(encoding="utf-8")), bool(row[1])
    except Exception as exc:
        raise ResultReadError(f"published result could not be read: {exc}") from exc


def _result_input_scope_id(result: dict[str, Any]) -> str:
    source = result["source"]
    return source.get("input_scope_id") or _input_scope_id(
        source["source_file_id"],
        source["sheet"],
        _normalize_selected_rows(source["selected_excel_rows"]),
    )


def _result_rule_id(result: dict[str, Any]) -> str:
    return result.get("aggregation_rule_id") or _sql_rule_id(result["applied_sql"])


def _result_transformation_rule_id(result: dict[str, Any]) -> str | None:
    return result.get("transformation_rule_id")


def list_published_results(
    database_path: Path,
    published_root: Path,
    limit: int = 20,
) -> list[dict[str, Any]]:
    database_path = database_path.resolve()
    if not database_path.exists():
        return []
    try:
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            rows = connection.execute(
                """
                SELECT pr.run_id
                FROM published_results AS pr
                JOIN pipeline_runs AS r USING (run_id)
                WHERE r.status = 'succeeded'
                ORDER BY pr.rowid DESC
                LIMIT ?
                """,
                [limit],
            ).fetchall()
        finally:
            connection.close()
    except Exception as exc:
        raise ResultReadError(f"result database could not be read: {exc}") from exc

    summaries: list[dict[str, Any]] = []
    for (run_id,) in rows:
        result, is_current = _load_published_result(
            database_path, published_root, run_id
        )
        summaries.append(
            {
                "run_id": run_id,
                "completed_at": result["run"]["completed_at"],
                "is_current": is_current,
                "source_file_id": result["source"]["source_file_id"],
                "sheet": result["source"]["sheet"],
                "selected_excel_rows": result["source"]["selected_excel_rows"],
                "input_scope_id": _result_input_scope_id(result),
                "aggregation_rule_id": _result_rule_id(result),
                "transformation_rule_id": _result_transformation_rule_id(result),
            }
        )
    return summaries


def _product_map(result: dict[str, Any]) -> dict[tuple[str, str], dict[str, Any]]:
    return {
        (item["observed_date"], item["stock_code"]): item
        for item in result["products"]
    }


def _source_locations(
    result: dict[str, Any], item: dict[str, Any] | None
) -> set[tuple[str, str, int]]:
    if item is None:
        return set()
    source_file_id = result["source"]["source_file_id"]
    return {
        (source_file_id, row["sheet"], int(row["source_row_number"]))
        for row in item["contributing_rows"]
    }


def _location_payload(location: tuple[str, str, int]) -> dict[str, Any]:
    return {
        "source_file_id": location[0],
        "sheet": location[1],
        "source_row_number": location[2],
    }


def compare_results(
    base: dict[str, Any], current: dict[str, Any]
) -> dict[str, Any]:
    base_source = base["source"]
    current_source = current["source"]
    source_changed = base_source["source_file_id"] != current_source["source_file_id"]
    sheet_changed = base_source["sheet"] != current_source["sheet"]
    rule_changed = _result_rule_id(base) != _result_rule_id(current)
    base_transformation_rule = _result_transformation_rule_id(base)
    current_transformation_rule = _result_transformation_rule_id(current)
    transformation_rule_unrecorded = (
        base_transformation_rule is None or current_transformation_rule is None
    )
    transformation_rule_changed = (
        base_transformation_rule != current_transformation_rule
    )
    scope_changed = _result_input_scope_id(base) != _result_input_scope_id(current)

    base_selected = set(base_source["selected_excel_rows"])
    current_selected = set(current_source["selected_excel_rows"])
    if not source_changed and not sheet_changed:
        selected_rows_added = sorted(current_selected - base_selected)
        selected_rows_removed = sorted(base_selected - current_selected)
    else:
        selected_rows_added = []
        selected_rows_removed = []

    rules_match = (
        not rule_changed
        and not transformation_rule_changed
        and not transformation_rule_unrecorded
    )
    if not source_changed and not sheet_changed and rules_match and scope_changed:
        change_code = "selection_scope_changed"
    elif not source_changed and not sheet_changed and rules_match and not scope_changed:
        change_code = "same_input_scope_and_rule"
    else:
        change_code = "source_sheet_or_rule_changed"

    base_products = _product_map(base)
    current_products = _product_map(current)
    groups: list[dict[str, Any]] = []
    for observed_date, stock_code in sorted(base_products.keys() | current_products.keys()):
        base_item = base_products.get((observed_date, stock_code))
        current_item = current_products.get((observed_date, stock_code))
        base_locations = _source_locations(base, base_item)
        current_locations = _source_locations(current, current_item)
        groups.append(
            {
                "observed_date": observed_date,
                "stock_code": stock_code,
                "base_state": "available" if base_item is not None else "not_present",
                "current_state": "available" if current_item is not None else "not_present",
                "base_sample_quantity_sum": (
                    base_item["sample_quantity_sum"] if base_item is not None else None
                ),
                "current_sample_quantity_sum": (
                    current_item["sample_quantity_sum"]
                    if current_item is not None
                    else None
                ),
                "base_observed_row_count": (
                    base_item["observed_row_count"] if base_item is not None else None
                ),
                "current_observed_row_count": (
                    current_item["observed_row_count"]
                    if current_item is not None
                    else None
                ),
                "added_source_rows": [
                    _location_payload(location)
                    for location in sorted(current_locations - base_locations)
                ],
                "removed_source_rows": [
                    _location_payload(location)
                    for location in sorted(base_locations - current_locations)
                ],
                "changed": (
                    (
                        None
                        if base_item is None
                        else (
                            base_item["sample_quantity_sum"],
                            base_item["observed_row_count"],
                        )
                    )
                    != (
                        None
                        if current_item is None
                        else (
                            current_item["sample_quantity_sum"],
                            current_item["observed_row_count"],
                        )
                    )
                    or base_locations != current_locations
                ),
            }
        )

    return {
        "comparison_state": "available",
        "base_run_id": base["run"]["run_id"],
        "current_run_id": current["run"]["run_id"],
        "change_basis": {
            "code": change_code,
            "source_file_changed": source_changed,
            "sheet_changed": sheet_changed,
            "selection_scope_changed": scope_changed,
            "aggregation_rule_changed": rule_changed,
            "transformation_rule_changed": transformation_rule_changed,
            "transformation_rule_unrecorded": transformation_rule_unrecorded,
            "selected_rows_added": selected_rows_added,
            "selected_rows_removed": selected_rows_removed,
        },
        "groups": groups,
    }


def compare_published_results(
    database_path: Path,
    published_root: Path,
    *,
    base_run_id: str,
    current_run_id: str,
) -> dict[str, Any]:
    base, _ = _load_published_result(database_path, published_root, base_run_id)
    current, _ = _load_published_result(
        database_path, published_root, current_run_id
    )
    return compare_results(base, current)


def load_run_history(database_path: Path, limit: int = 20) -> list[dict[str, Any]]:
    database_path = database_path.resolve()
    if not database_path.exists():
        return []
    try:
        connection = duckdb.connect(str(database_path), read_only=True)
        try:
            rows = connection.execute(
                """
                SELECT r.run_id, r.source_file_id, r.status, r.started_at,
                       r.completed_at, r.error_message, s.selected_excel_rows,
                       r.owner_pid, r.owner_token, r.recovered_at,
                       r.recovery_reason, r.recovered_artifact_path
                FROM pipeline_runs AS r
                LEFT JOIN run_input_scopes AS s USING (run_id)
                ORDER BY r.rowid DESC
                LIMIT ?
                """,
                [limit],
            ).fetchall()
        finally:
            connection.close()
    except Exception as exc:
        raise ResultReadError(f"run history could not be read: {exc}") from exc

    return [
        {
            "run_id": row[0],
            "source_file_id": row[1],
            "status": row[2],
            "started_at": row[3].isoformat(timespec="seconds"),
            "completed_at": row[4].isoformat(timespec="seconds") if row[4] else None,
            "error_message": row[5],
            "selected_excel_rows": json.loads(row[6]) if row[6] else None,
            "owner_pid": row[7],
            "owner_token": row[8],
            "recovered_at": row[9].isoformat(timespec="seconds") if row[9] else None,
            "recovery_reason": row[10],
            "recovered_artifact_path": row[11],
        }
        for row in rows
    ]
