from __future__ import annotations

import hashlib
import json
import tempfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import load_workbook

from .config import SOURCE_MEANING_SQL
from .pipeline import PipelineFailure, load_manifest, run_pipeline


DEFAULT_SOURCE_MEANING_ROWS = (2, 143, 156)
SUPPORTED_INVOICE_WHITESPACE = " \t\r\n"


def _rule_id(sql_text: str) -> str:
    return f"sha256:{hashlib.sha256(sql_text.encode('utf-8')).hexdigest()}"


def _independent_marker(invoice_no: Any) -> str:
    if invoice_no is None or isinstance(invoice_no, bool):
        return "unknown"
    if isinstance(invoice_no, str):
        observed = invoice_no.strip(SUPPORTED_INVOICE_WHITESPACE)
    elif isinstance(invoice_no, int):
        observed = str(invoice_no)
    elif isinstance(invoice_no, float) and invoice_no.is_integer():
        observed = str(int(invoice_no))
    else:
        return "unknown"
    if not observed:
        return "unknown"
    return "true" if observed[0].upper() == "C" else "false"


def _lossless_integer(value: Any, row_number: int) -> int:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError) as exc:
        raise PipelineFailure(
            f"independent check found invalid Quantity at Excel row {row_number}"
        ) from exc
    if not number.is_finite() or number != number.to_integral_value():
        raise PipelineFailure(
            f"independent check found non-integer Quantity at Excel row {row_number}"
        )
    return int(number)


def _group_payload(
    key: tuple[str, str, str | None, str],
    value: dict[str, Any],
) -> dict[str, Any]:
    return {
        "observed_date": key[0],
        "stock_code": key[1],
        "description": key[2],
        "source_cancellation_marker": key[3],
        "quantity_sum": value["quantity_sum"],
        "observed_row_count": value["observed_row_count"],
        "contributing_excel_rows": value["contributing_excel_rows"],
    }


def independent_workbook_observation(
    source_path: Path,
    sheet_name: str,
    selected_rows: list[int],
) -> list[dict[str, Any]]:
    """Calculate the same small observation directly with openpyxl and Python.

    This path does not read the DuckDB view or the observation SQL. It is an
    independent comparison for the selected source rows, not a second product
    policy implementation.
    """

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    wanted = set(selected_rows)
    found: set[int] = set()
    groups: dict[tuple[str, str, str | None, str], dict[str, Any]] = defaultdict(
        lambda: {
            "quantity_sum": 0,
            "observed_row_count": 0,
            "contributing_excel_rows": [],
        }
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise PipelineFailure(f"workbook sheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                max_row=max(selected_rows),
                max_col=6,
                values_only=True,
            ),
            start=2,
        ):
            if row_number not in wanted:
                continue
            found.add(row_number)
            invoice_no, stock_code, description, quantity, invoice_date, _ = row
            if not isinstance(invoice_date, (date, datetime)):
                raise PipelineFailure(
                    f"independent check found invalid InvoiceDate at Excel row {row_number}"
                )
            observed_date = (
                invoice_date.date().isoformat()
                if isinstance(invoice_date, datetime)
                else invoice_date.isoformat()
            )
            key = (
                observed_date,
                str(stock_code),
                None if description is None else str(description),
                _independent_marker(invoice_no),
            )
            group = groups[key]
            group["quantity_sum"] += _lossless_integer(quantity, row_number)
            group["observed_row_count"] += 1
            group["contributing_excel_rows"].append(row_number)
    finally:
        workbook.close()

    missing = sorted(wanted - found)
    if missing:
        raise PipelineFailure(
            f"independent check did not find selected Excel rows: {missing}"
        )
    ordered_keys = sorted(
        groups,
        key=lambda key: (
            key[0],
            key[1],
            key[2] is None,
            key[2] or "",
            key[3],
        ),
    )
    return [_group_payload(key, groups[key]) for key in ordered_keys]


def _sql_observation(
    database_path: Path,
    run_id: str,
    observation_sql: str,
) -> list[dict[str, Any]]:
    connection = duckdb.connect(str(database_path), read_only=True)
    try:
        rows = connection.execute(observation_sql, [run_id]).fetchall()
    finally:
        connection.close()
    return [
        {
            "observed_date": row[0].isoformat(),
            "stock_code": row[1],
            "description": row[2],
            "source_cancellation_marker": row[3],
            "quantity_sum": int(row[4]),
            "observed_row_count": int(row[5]),
            "contributing_excel_rows": [int(value) for value in row[6]],
        }
        for row in rows
    ]


def observe_source_meaning(
    *,
    source_path: Path,
    manifest_path: Path,
    selected_rows: list[int] | None = None,
) -> dict[str, Any]:
    selected_rows = list(
        DEFAULT_SOURCE_MEANING_ROWS if selected_rows is None else selected_rows
    )
    manifest = load_manifest(manifest_path)
    observation_sql = SOURCE_MEANING_SQL.read_text(encoding="utf-8").strip()

    with tempfile.TemporaryDirectory(prefix="shopping-source-meaning-") as temporary:
        temporary_root = Path(temporary)
        database_path = temporary_root / "observation.duckdb"
        pipeline_result = run_pipeline(
            source_path=source_path,
            manifest_path=manifest_path,
            database_path=database_path,
            published_root=temporary_root / "published",
            selected_rows=selected_rows,
        )
        sql_groups = _sql_observation(
            database_path,
            pipeline_result["run"]["run_id"],
            observation_sql,
        )

    independent_groups = independent_workbook_observation(
        source_path,
        manifest["sheet"],
        pipeline_result["source"]["selected_excel_rows"],
    )
    if sql_groups != independent_groups:
        raise PipelineFailure(
            "source-meaning SQL did not match the independent openpyxl calculation: "
            + json.dumps(
                {"sql": sql_groups, "independent": independent_groups},
                ensure_ascii=False,
                sort_keys=True,
            )
        )

    return {
        "result_state": "observed_for_selected_source_rows",
        "source": {
            "dataset": manifest["dataset"],
            "source_url": manifest["source_url"],
            "source_file_name": source_path.name,
            "sha256": pipeline_result["source"]["sha256"],
            "source_file_id": pipeline_result["source"]["source_file_id"],
            "sheet": manifest["sheet"],
            "selected_excel_rows": pipeline_result["source"][
                "selected_excel_rows"
            ],
            "input_scope_id": pipeline_result["source"]["input_scope_id"],
        },
        "execution_evidence": {
            "isolated_run_id": pipeline_result["run"]["run_id"],
            "pipeline_result_state": pipeline_result["result_state"],
            "completed_at": pipeline_result["run"]["completed_at"],
            "storage": (
                "Temporary DuckDB and result artifact were isolated from the default "
                "database and removed after this command produced its output."
            ),
        },
        "rules": {
            "source_transformation_rule_id": pipeline_result[
                "transformation_rule_id"
            ],
            "observation_rule_id": _rule_id(observation_sql),
            "applied_source_transformation_sql": pipeline_result[
                "applied_transformation_sql"
            ],
            "applied_observation_sql": observation_sql,
        },
        "observation_contract": {
            "grouping": (
                "Source-displayed InvoiceDate date part, StockCode, raw Description, "
                "and the derived source cancellation marker."
            ),
            "quantity_sum": "Signed Quantity summed only inside each observed group.",
            "observed_row_count": "Selected source rows present in that group.",
            "zero_and_absence": (
                "A present group may have quantity_sum 0 while observed_row_count is "
                "positive. A group with no observed rows is not emitted and is not "
                "turned into a zero-valued result."
            ),
            "not_claimed": [
                "sales, product cancellation, discount adjustment, refund, payment, or net revenue",
                "original transaction linkage or original-sale attribution date",
                "whole-day or full-dataset distribution",
            ],
        },
        "source_cancellation_marker_contract": pipeline_result[
            "source_cancellation_marker_contract"
        ],
        "groups": sql_groups,
        "independent_check": {
            "engine": "openpyxl plus Python; no DuckDB view or observation SQL",
            "matches": True,
            "groups": independent_groups,
        },
        "original_sale_date": {
            "state": "not_calculable_from_observed_fields",
            "reason": (
                "The source InvoiceDate describes the observed transaction creation "
                "time. The selected fields and current contract do not identify the "
                "original transaction or define which date should receive an adjustment."
            ),
            "needed_evidence": [
                "a verified original-transaction link or business key",
                "the business classification of rows such as D / Discount",
                "an accepted attribution-date rule",
            ],
        },
    }
