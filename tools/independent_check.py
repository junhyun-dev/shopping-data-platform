#!/usr/bin/env python3
"""Independent workbook-only check; it does not import production pipeline code."""

import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "config" / "source.json"


def main() -> None:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    workbook_path = ROOT / manifest["local_file"]
    selected = set(manifest["selected_excel_rows"])
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[manifest["sheet"]]
        totals: dict[tuple[str, str], dict[str, int]] = defaultdict(
            lambda: {"sample_quantity_sum": 0, "observed_row_count": 0}
        )
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_row=max(selected), max_col=8, values_only=True),
            start=2,
        ):
            if row_number not in selected:
                continue
            invoice_date = row[4]
            if not isinstance(invoice_date, datetime):
                raise ValueError(f"Excel row {row_number} has no datetime InvoiceDate")
            key = (str(row[1]), invoice_date.date().isoformat())
            totals[key]["sample_quantity_sum"] += int(row[3])
            totals[key]["observed_row_count"] += 1
    finally:
        workbook.close()

    result = [
        {"stock_code": key[0], "observed_date": key[1], **value}
        for key, value in sorted(totals.items())
    ]
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
