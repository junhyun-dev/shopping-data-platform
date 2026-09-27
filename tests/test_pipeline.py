from __future__ import annotations

import hashlib
import json
import os
import select
import signal
import subprocess
import sys
from copy import deepcopy
from datetime import datetime
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

import shopping_data.pipeline as pipeline_module
from shopping_data.config import CONTRIBUTIONS_SQL
from shopping_data.pipeline import (
    PROCESS_CHECKPOINTS,
    PipelineBusy,
    PipelineFailure,
    ResultNotFound,
    ResultReadError,
    compare_published_results,
    compare_results,
    load_current_result,
    load_published_result,
    list_published_results,
    load_run_history,
    run_pipeline,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


HEADERS = [
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
]

# Adapted from UCI Online Retail (Chen, 2015), CC BY 4.0; see README.
# Customer IDs are replaced with the fictional TEST-CUSTOMER value.
# Additional rows exercise synthetic edge cases; this is not a source extract.
ROWS = {
    2: [536365, "85123A", "WHITE HANGING HEART T-LIGHT HOLDER", 6, datetime(2010, 12, 1, 8, 26), 2.55, "TEST-CUSTOMER", "United Kingdom"],
    3: [536365, 71053, "WHITE METAL LANTERN", 6, datetime(2010, 12, 1, 8, 26), 3.39, "TEST-CUSTOMER", "United Kingdom"],
    4: [536365, "84406B", "CREAM CUPID HEARTS COAT HANGER", 8, datetime(2010, 12, 1, 8, 26), 2.75, "TEST-CUSTOMER", "United Kingdom"],
    5: [536365, "84029G", "KNITTED UNION FLAG HOT WATER BOTTLE", 6, datetime(2010, 12, 1, 8, 26), 3.39, "TEST-CUSTOMER", "United Kingdom"],
    6: [536365, "84029E", "RED WOOLLY HOTTIE WHITE HEART.", 6, datetime(2010, 12, 1, 8, 26), 3.39, "TEST-CUSTOMER", "United Kingdom"],
    7: [536365, 22752, "SET 7 BABUSHKA NESTING BOXES", 2, datetime(2010, 12, 1, 8, 26), 7.65, "TEST-CUSTOMER", "United Kingdom"],
    8: [536365, 21730, "GLASS STAR FROSTED T-LIGHT HOLDER", 6, datetime(2010, 12, 1, 8, 26), 4.25, "TEST-CUSTOMER", "United Kingdom"],
    9: [536365, "ZERO", "ZERO QUANTITY TEST ROW", 0, datetime(2010, 12, 1, 8, 26), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    10: [536365, "85123A", "WHITE HANGING HEART T-LIGHT HOLDER", 6, datetime(2010, 12, 1, 8, 26), 2.55, "TEST-CUSTOMER", "United Kingdom"],
    11: [None, "MISSING", "MISSING INVOICE NUMBER", 1, datetime(2010, 12, 1, 9, 0), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    12: [" c-test", "LOWER-C", "LOWERCASE PREFIX", 1, datetime(2010, 12, 1, 9, 1), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    13: ["   ", "BLANK", "BLANK INVOICE NUMBER", 1, datetime(2010, 12, 1, 9, 2), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    14: [True, "UNSUPPORTED", "UNSUPPORTED INVOICE TYPE", 1, datetime(2010, 12, 1, 9, 3), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    15: ["\t", "TAB-BLANK", "TAB-ONLY INVOICE NUMBER", 1, datetime(2010, 12, 1, 9, 4), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    16: ["\tC123\n", "TAB-C", "C PREFIX AFTER TAB", 1, datetime(2010, 12, 1, 9, 5), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    17: [536400, "MIX", "MIX DESCRIPTION A", 5, datetime(2010, 12, 1, 10, 0), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    18: ["C536401", "MIX", "MIX DESCRIPTION B", -5, datetime(2010, 12, 1, 10, 1), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    19: [None, "MIX", "MIX DESCRIPTION C", 2, datetime(2010, 12, 1, 10, 2), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    20: ["C536402", "MIX", "MIX DESCRIPTION D", -2, datetime(2010, 12, 1, 10, 3), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    21: [536403, "BIG", "EXACT INTEGER TEXT", "9007199254740993", datetime(2010, 12, 1, 10, 4), 1.0, "TEST-CUSTOMER", "United Kingdom"],
    143: ["C536379", "D", "Discount", -1, datetime(2010, 12, 1, 9, 41), 27.5, "TEST-CUSTOMER", "United Kingdom"],
    156: ["C536383", "35004C", "SET OF 3 COLOURED  FLYING DUCKS", -1, datetime(2010, 12, 1, 9, 49), 4.65, "TEST-CUSTOMER", "United Kingdom"],
    202: [536389, "35004C", "SET OF 3 COLOURED  FLYING DUCKS", 6, datetime(2010, 12, 1, 10, 3), 4.65, "TEST-CUSTOMER", "United Kingdom"],
    299: [536397, "35004C", "SET OF 3 COLOURED  FLYING DUCKS", 48, datetime(2010, 12, 1, 10, 51), 4.65, "TEST-CUSTOMER", "United Kingdom"],
}


def write_workbook(path: Path, selected_rows: list[int]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Online Retail"
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=1, column=column, value=header)
    for row_number in selected_rows:
        for column, value in enumerate(ROWS[row_number], start=1):
            sheet.cell(row=row_number, column=column, value=value)
    workbook.save(path)


def write_manifest(path: Path, workbook_path: Path, selected_rows: list[int]) -> None:
    digest = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    path.write_text(
        json.dumps(
            {
                "dataset": "UCI-derived test workbook with synthetic edge cases",
                "source_url": "local-test-only",
                "expected_sha256": digest,
                "sheet": "Online Retail",
                "selected_excel_rows": selected_rows,
                "scope_note": "Synthetic test rows only.",
            }
        ),
        encoding="utf-8",
    )


def local_paths(tmp_path: Path, selected_rows: list[int] | None = None):
    selected_rows = selected_rows or [2, 3, 4, 5, 6, 7, 8, 143]
    workbook = tmp_path / "sample.xlsx"
    manifest = tmp_path / "manifest.json"
    database = tmp_path / "shopping.duckdb"
    published = tmp_path / "published"
    write_workbook(workbook, selected_rows)
    write_manifest(manifest, workbook, selected_rows)
    return workbook, manifest, database, published


def source_stage_record(
    *,
    row_number: int,
    quantity_raw: object = "1",
    invoice_no_source_kind: object = "text",
) -> dict[str, object]:
    return {
        "source_file_id": "file-sha",
        "sheet_name": "Online Retail",
        "source_row_number": row_number,
        "invoice_no_raw": None,
        "invoice_no_source_kind": invoice_no_source_kind,
        "stock_code_raw": '상품-\\"A',
        "description_raw": '한글, 따옴표 ", 역슬래시 \\\\, 줄바꿈\n탭\t',
        "quantity_raw": quantity_raw,
        "invoice_date_raw": "2010-12-01 08:26:00",
        "unit_price_raw": "0001.234567890123456789",
        "customer_id_raw": "null",
        "country_raw": "",
        "loaded_at": "2026-09-26T12:34:56+00:00",
    }


def start_crash_worker(
    *,
    workbook: Path,
    manifest: Path,
    database: Path,
    published: Path,
    selected_rows: list[int],
    checkpoint: str,
) -> tuple[subprocess.Popen[str], dict]:
    env = os.environ.copy()
    existing_pythonpath = env.get("PYTHONPATH")
    env["PYTHONPATH"] = (
        str(PROJECT_ROOT)
        if not existing_pythonpath
        else f"{PROJECT_ROOT}{os.pathsep}{existing_pythonpath}"
    )
    process = subprocess.Popen(
        [
            sys.executable,
            str(PROJECT_ROOT / "tests" / "crash_worker.py"),
            "--source",
            str(workbook),
            "--manifest",
            str(manifest),
            "--database",
            str(database),
            "--published-root",
            str(published),
            "--rows",
            *[str(row) for row in selected_rows],
            "--checkpoint",
            checkpoint,
        ],
        cwd=PROJECT_ROOT,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    assert process.stdout is not None
    ready, _, _ = select.select([process.stdout], [], [], 20)
    if not ready:
        process.kill()
        stdout, stderr = process.communicate(timeout=10)
        raise AssertionError(
            f"child did not reach checkpoint {checkpoint}; stdout={stdout!r}, stderr={stderr!r}"
        )
    payload = json.loads(process.stdout.readline())
    assert payload["state"] == "ready_for_forced_termination"
    assert payload["checkpoint"] == checkpoint
    assert process.poll() is None
    return process, payload


def kill_owned_worker(process: subprocess.Popen[str]) -> None:
    if process.poll() is None:
        process.kill()
    process.wait(timeout=10)
    assert process.returncode == -signal.SIGKILL


def product(result: dict, stock_code: str) -> dict:
    return next(item for item in result["products"] if item["stock_code"] == stock_code)


def test_json_source_stage_preserves_raw_text_and_rejects_invalid_shape() -> None:
    record = source_stage_record(
        row_number=2,
        quantity_raw="9007199254740993123456789",
    )
    connection = duckdb.connect(":memory:")
    try:
        pipeline_module._stage_source_rows_from_json(
            connection,
            payload=json.dumps([record], ensure_ascii=False),
            source_file_id="file-sha",
            sheet_name="Online Retail",
            selected_rows=[2],
        )
        preserved = connection.execute(
            """
            SELECT invoice_no_raw, stock_code_raw, description_raw, quantity_raw,
                   invoice_date_raw, unit_price_raw, customer_id_raw, country_raw
            FROM pipeline_source_stage
            """
        ).fetchone()
        assert preserved == (
            record["invoice_no_raw"],
            record["stock_code_raw"],
            record["description_raw"],
            record["quantity_raw"],
            record["invoice_date_raw"],
            record["unit_price_raw"],
            record["customer_id_raw"],
            record["country_raw"],
        )
    finally:
        connection.close()

    missing = dict(record)
    missing.pop("country_raw")
    extra = dict(record, unexpected="x")
    numeric_raw = dict(record, quantity_raw=7)
    replaced_location = dict(record, source_row_number=99)
    cases = [
        ([missing], [2], "missing or extra fields"),
        ([extra], [2], "missing or extra fields"),
        ([numeric_raw], [2], "invalid field types"),
        ([record], [2, 3], "row count mismatch"),
        ([replaced_location], [2], "staged source locations"),
    ]
    for records, selected_rows, message in cases:
        invalid_connection = duckdb.connect(":memory:")
        try:
            with pytest.raises(PipelineFailure, match=message):
                pipeline_module._stage_source_rows_from_json(
                    invalid_connection,
                    payload=json.dumps(records, ensure_ascii=False),
                    source_file_id="file-sha",
                    sheet_name="Online Retail",
                    selected_rows=selected_rows,
                )
        finally:
            invalid_connection.close()

    for payload, message in [
        (json.dumps({"row": record}, ensure_ascii=False), "must be a JSON array"),
        ("not-json", "could not be staged"),
    ]:
        invalid_connection = duckdb.connect(":memory:")
        try:
            with pytest.raises(PipelineFailure, match=message):
                pipeline_module._stage_source_rows_from_json(
                    invalid_connection,
                    payload=payload,
                    source_file_id="file-sha",
                    sheet_name="Online Retail",
                    selected_rows=[2],
                )
        finally:
            invalid_connection.close()


def test_json_source_stage_materializes_scalar_items_once_before_validation() -> None:
    records = [
        dict(
            source_stage_record(
                row_number=row_number,
                quantity_raw=str(
                    9_007_199_254_740_993
                    if row_number % 101 == 0
                    else row_number - 500
                ),
            ),
            stock_code_raw=f"SKU-{row_number % 97:02d}",
            unit_price_raw=None if row_number % 11 == 0 else "1.234567",
        )
        for row_number in range(2, 519)
    ]
    selected_rows = list(range(2, 519))
    connection = duckdb.connect(":memory:")
    queries: list[str] = []

    class RecordingConnection:
        def execute(
            self, query: str, parameters: object | None = None
        ) -> duckdb.DuckDBPyConnection:
            queries.append(" ".join(query.split()))
            if parameters is None:
                return connection.execute(query)
            return connection.execute(query, parameters)

    try:
        connection.execute("SET memory_limit='512MB'")
        connection.execute("SET threads=1")
        pipeline_module._stage_source_rows_from_json(
            RecordingConnection(),
            payload=json.dumps(records, ensure_ascii=False),
            source_file_id="file-sha",
            sheet_name="Online Retail",
            selected_rows=selected_rows,
        )
        assert connection.execute(
            "SELECT COUNT(*) FROM pipeline_source_items"
        ).fetchone()[0] == 517
        assert connection.execute(
            "SELECT COUNT(*) FROM pipeline_source_stage"
        ).fetchone()[0] == 517
        assert connection.execute(
            """
            SELECT quantity_raw, unit_price_raw
            FROM pipeline_source_stage
            WHERE source_row_number = 101
            """
        ).fetchone() == ("9007199254740993", "1.234567")

        json_each_queries = [query for query in queries if "json_each(" in query]
        assert len(json_each_queries) == 1
        assert (
            "FROM json_each((SELECT document FROM pipeline_source_payload))"
            in json_each_queries[0]
        )
        assert not any(
            "pipeline_source_payload, json_each(document)" in query
            for query in queries
        )
        assert sum("FROM pipeline_source_items" in query for query in queries) == 4
    finally:
        connection.close()


def test_staged_write_preserves_first_raw_values_fills_only_null_kind_and_rolls_back() -> None:
    connection = duckdb.connect(":memory:")
    try:
        pipeline_module.initialize_database(connection)
        now = pipeline_module.utc_now()
        connection.execute(
            """
            INSERT INTO pipeline_runs
                (run_id, source_file_id, input_path, status, started_at)
            VALUES ('run-json', 'file-sha', 'test.xlsx', 'running', ?)
            """,
            [now],
        )
        connection.execute(
            """
            INSERT INTO source_rows VALUES
                ('file-sha', 'Online Retail', 2, 'old-2', 'text', 'P2', 'old',
                 '1', '2010-12-01 08:26:00', '1', NULL, 'UK', ?),
                ('file-sha', 'Online Retail', 3, 'old-3', NULL, 'P3', 'old',
                 '2', '2010-12-01 08:27:00', '2', NULL, 'UK', ?)
            """,
            [now, now],
        )
        records = [
            dict(
                source_stage_record(
                    row_number=2,
                    quantity_raw="99",
                    invoice_no_source_kind="integer",
                ),
                stock_code_raw="P2",
            ),
            dict(
                source_stage_record(
                    row_number=3,
                    quantity_raw="88",
                    invoice_no_source_kind="integer",
                ),
                stock_code_raw="P3",
            ),
            dict(
                source_stage_record(
                    row_number=4,
                    quantity_raw="7",
                    invoice_no_source_kind="text",
                ),
                stock_code_raw="P4",
            ),
        ]

        connection.execute("BEGIN TRANSACTION")
        pipeline_module._stage_source_rows_from_json(
            connection,
            payload=json.dumps(records, ensure_ascii=False),
            source_file_id="file-sha",
            sheet_name="Online Retail",
            selected_rows=[2, 3, 4],
        )
        pipeline_module._persist_staged_source_rows(connection, run_id="run-json")
        assert connection.execute(
            """
            SELECT source_row_number, quantity_raw, invoice_no_source_kind
            FROM source_rows ORDER BY source_row_number
            """
        ).fetchall() == [(2, "1", "text"), (3, "2", "integer"), (4, "7", "text")]
        assert connection.execute(
            "SELECT COUNT(*) FROM run_input_rows WHERE run_id = 'run-json'"
        ).fetchone()[0] == 3
        connection.execute("ROLLBACK")

        assert connection.execute(
            """
            SELECT source_row_number, quantity_raw, invoice_no_source_kind
            FROM source_rows ORDER BY source_row_number
            """
        ).fetchall() == [(2, "1", "text"), (3, "2", None)]
        assert connection.execute(
            "SELECT COUNT(*) FROM run_input_rows WHERE run_id = 'run-json'"
        ).fetchone()[0] == 0
    finally:
        connection.close()


def test_json_stage_supports_legacy_additive_source_rows_column_order(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    connection = duckdb.connect(str(database))
    try:
        connection.execute(
            """
            CREATE TABLE source_rows (
                source_file_id VARCHAR NOT NULL,
                sheet_name VARCHAR NOT NULL,
                source_row_number BIGINT NOT NULL,
                invoice_no_raw VARCHAR,
                stock_code_raw VARCHAR,
                description_raw VARCHAR,
                quantity_raw VARCHAR,
                invoice_date_raw VARCHAR,
                unit_price_raw VARCHAR,
                customer_id_raw VARCHAR,
                country_raw VARCHAR,
                loaded_at TIMESTAMPTZ NOT NULL,
                PRIMARY KEY (source_file_id, sheet_name, source_row_number)
            )
            """
        )
    finally:
        connection.close()

    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )

    connection = duckdb.connect(str(database))
    try:
        column_names = [
            row[1]
            for row in connection.execute(
                "PRAGMA table_info('source_rows')"
            ).fetchall()
        ]
        assert column_names[-2:] == ["loaded_at", "invoice_no_source_kind"]
        stored_rows = connection.execute(
            """
            SELECT source_row_number, invoice_no_raw, stock_code_raw,
                   description_raw, quantity_raw, invoice_date_raw,
                   unit_price_raw, customer_id_raw, country_raw,
                   invoice_no_source_kind
            FROM source_rows
            ORDER BY source_row_number
            """
        ).fetchall()
        assert stored_rows == [
            (
                2,
                "536365",
                "85123A",
                "WHITE HANGING HEART T-LIGHT HOLDER",
                "6",
                "2010-12-01 08:26:00",
                "2.55",
                "TEST-CUSTOMER",
                "United Kingdom",
                "integer",
            ),
            (
                3,
                "536365",
                "71053",
                "WHITE METAL LANTERN",
                "6",
                "2010-12-01 08:26:00",
                "3.39",
                "TEST-CUSTOMER",
                "United Kingdom",
                "integer",
            ),
        ]
        connection.execute(
            """
            UPDATE source_rows
            SET invoice_no_source_kind = NULL
            WHERE source_row_number = 2
            """
        )
    finally:
        connection.close()

    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    assert second["source"]["input_scope_id"] == first["source"]["input_scope_id"]
    assert second["products"] == first["products"]

    connection = duckdb.connect(str(database), read_only=True)
    try:
        assert connection.execute(
            """
            SELECT invoice_no_source_kind
            FROM source_rows
            WHERE source_row_number = 2
            """
        ).fetchone()[0] == "integer"
        assert connection.execute("SELECT COUNT(*) FROM source_rows").fetchone()[0] == 2
        assert connection.execute("SELECT COUNT(*) FROM run_input_rows").fetchone()[0] == 4
    finally:
        connection.close()

    assert load_current_result(database, published)["run"]["run_id"] == second[
        "run"
    ]["run_id"]


def marker_breakdown(product_result: dict) -> dict[str, dict]:
    return {
        item["source_cancellation_marker"]: item
        for item in product_result["source_cancellation_marker_breakdown"]["groups"]
    }


class RecordingConnection:
    def __init__(self, connection: duckdb.DuckDBPyConnection) -> None:
        self.connection = connection
        self.executions: list[tuple[str, object]] = []

    def execute(self, query: str, parameters: object = None):
        self.executions.append((query, parameters))
        if parameters is None:
            return self.connection.execute(query)
        return self.connection.execute(query, parameters)


def test_result_build_reads_all_run_contributions_once_without_cross_run_rows(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(
        tmp_path, [2, 3, 10, 143]
    )
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[10, 143],
    )

    connection = duckdb.connect(str(database), read_only=True)
    recording = RecordingConnection(connection)
    try:
        rebuilt = pipeline_module._build_result(
            recording,
            run_id=first["run"]["run_id"],
            source_file_id=first["source"]["source_file_id"],
            source_sha256=first["source"]["sha256"],
            source_path=workbook,
            manifest=json.loads(manifest.read_text(encoding="utf-8")),
            selected_rows=first["source"]["selected_excel_rows"],
            input_scope_id=first["source"]["input_scope_id"],
            aggregation_rule_id=first["aggregation_rule_id"],
            transformation_rule_id=first["transformation_rule_id"],
            completed_at=datetime.fromisoformat(first["run"]["completed_at"]),
        )
    finally:
        connection.close()

    contribution_sql = CONTRIBUTIONS_SQL.read_text(encoding="utf-8").strip()
    contribution_query_count = sum(
        query.strip() == contribution_sql for query, _ in recording.executions
    )
    assert contribution_query_count == 1
    assert rebuilt == first


def test_bulk_contribution_reconciliation_failure_keeps_last_good_result(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    last_good = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    broken_sql = tmp_path / "broken_contribution_rows.sql"
    broken_sql.write_text(
        CONTRIBUTIONS_SQL.read_text(encoding="utf-8").replace(
            "WHERE i.run_id = ?",
            "WHERE i.run_id = ? AND t.stock_code <> '71053'",
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(pipeline_module, "CONTRIBUTIONS_SQL", broken_sql)

    with pytest.raises(
        PipelineFailure,
        match="source cancellation marker breakdown does not reconcile",
    ):
        run_pipeline(
            source_path=workbook,
            manifest_path=manifest,
            database_path=database,
            published_root=published,
        )

    assert load_current_result(database, published) == last_good
    history = load_run_history(database)
    assert history[0]["status"] == "failed"
    assert history[1]["run_id"] == last_good["run"]["run_id"]


def test_repeated_input_is_idempotent_and_products_do_not_mix(tmp_path: Path) -> None:
    selected = [2, 3, 4, 5, 6, 7, 8, 10, 143]
    workbook, manifest, database, published = local_paths(tmp_path, selected)
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    first_85123a = product(first, "85123A")
    second_85123a = product(second, "85123A")
    assert first_85123a["sample_quantity_sum"] == 12
    assert second_85123a["sample_quantity_sum"] == 12
    assert second_85123a["observed_row_count"] == 2
    assert [row["source_row_number"] for row in second_85123a["contributing_rows"]] == [2, 10]

    numeric_code = product(second, "71053")
    assert numeric_code["sample_quantity_sum"] == 6
    assert [row["source_row_number"] for row in numeric_code["contributing_rows"]] == [3]

    connection = duckdb.connect(str(database), read_only=True)
    try:
        assert connection.execute("SELECT COUNT(*) FROM source_files").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM source_rows").fetchone()[0] == 9
        assert connection.execute(
            "SELECT COUNT(*) FROM published_results WHERE is_current = TRUE"
        ).fetchone()[0] == 1
    finally:
        connection.close()

    current = load_current_result(database, published)
    assert current["run"]["run_id"] == second["run"]["run_id"]
    assert "customer" not in json.dumps(current).lower()


@pytest.mark.parametrize("failure_point", ["after_data", "after_artifact"])
def test_incomplete_run_does_not_replace_last_good_result(
    tmp_path: Path, failure_point: str
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path)
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    with pytest.raises(PipelineFailure, match="simulated failure"):
        run_pipeline(
            source_path=workbook,
            manifest_path=manifest,
            database_path=database,
            published_root=published,
            failure_point=failure_point,
        )

    current = load_current_result(database, published)
    assert current["run"]["run_id"] == first["run"]["run_id"]
    history = load_run_history(database)
    assert history[0]["status"] == "failed"
    assert history[1]["status"] == "succeeded"


@pytest.mark.parametrize("mutation", ["extra", "replace"])
def test_invalid_staged_locations_roll_back_without_replacing_last_good(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mutation: str,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    last_good = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    original_builder = pipeline_module._build_source_stage_payload

    def malformed_payload(**kwargs: object) -> str:
        records = json.loads(original_builder(**kwargs))
        if mutation == "extra":
            added = dict(records[0], source_row_number=99)
            records.append(added)
        else:
            records[0] = dict(records[0], source_row_number=99)
        return json.dumps(records, ensure_ascii=False, separators=(",", ":"))

    monkeypatch.setattr(
        pipeline_module, "_build_source_stage_payload", malformed_payload
    )
    expected_error = (
        "row count mismatch" if mutation == "extra" else "staged source locations"
    )
    with pytest.raises(PipelineFailure, match=expected_error):
        run_pipeline(
            source_path=workbook,
            manifest_path=manifest,
            database_path=database,
            published_root=published,
        )

    assert load_current_result(database, published) == last_good
    history = load_run_history(database)
    assert history[0]["status"] == "failed"
    assert history[1]["run_id"] == last_good["run"]["run_id"]
    connection = duckdb.connect(str(database), read_only=True)
    try:
        assert connection.execute("SELECT COUNT(*) FROM source_rows").fetchone()[0] == 2
        assert connection.execute(
            "SELECT COUNT(*) FROM run_input_rows WHERE run_id = ?",
            [history[0]["run_id"]],
        ).fetchone()[0] == 0
        assert connection.execute(
            "SELECT COUNT(*) FROM published_results"
        ).fetchone()[0] == 1
    finally:
        connection.close()


def test_unreadable_input_does_not_replace_last_good_result(tmp_path: Path) -> None:
    workbook, manifest, database, published = local_paths(tmp_path)
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    unreadable = tmp_path / "not-a-workbook.xlsx"
    unreadable.write_bytes(b"not an xlsx file")
    bad_manifest = tmp_path / "bad-manifest.json"
    write_manifest(bad_manifest, unreadable, [2])

    with pytest.raises(PipelineFailure, match="workbook could not be read"):
        run_pipeline(
            source_path=unreadable,
            manifest_path=bad_manifest,
            database_path=database,
            published_root=published,
        )

    assert load_current_result(database, published)["run"]["run_id"] == first["run"]["run_id"]


def test_fractional_quantity_is_rejected_without_replacing_last_good_result(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path)
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    fractional = tmp_path / "fractional.xlsx"
    write_workbook(fractional, [2])
    changed = load_workbook(fractional)
    changed["Online Retail"]["D2"] = 1.5
    changed["Online Retail"]["F2"] = 1.234567
    changed.save(fractional)
    changed.close()
    fractional_manifest = tmp_path / "fractional-manifest.json"
    write_manifest(fractional_manifest, fractional, [2])

    with pytest.raises(PipelineFailure, match="quantity cannot be represented losslessly"):
        run_pipeline(
            source_path=fractional,
            manifest_path=fractional_manifest,
            database_path=database,
            published_root=published,
        )

    assert load_current_result(database, published)["run"]["run_id"] == first["run"]["run_id"]
    history = load_run_history(database)
    assert history[0]["status"] == "failed"
    assert "Excel rows: 2" in history[0]["error_message"]


def test_unit_price_source_text_is_not_truncated(tmp_path: Path) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2])
    changed = load_workbook(workbook)
    changed["Online Retail"]["D2"] = 1
    changed["Online Retail"]["F2"] = 1.234567
    changed.save(workbook)
    changed.close()
    write_manifest(manifest, workbook, [2])

    result = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    row = product(result, "85123A")["contributing_rows"][0]
    assert row["unit_price_source_text"] == "1.234567"


def test_zero_quantity_is_present_not_missing(tmp_path: Path) -> None:
    selected = [2, 9, 143]
    workbook, manifest, database, published = local_paths(tmp_path, selected)
    result = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    zero = product(result, "ZERO")
    assert zero["data_state"] == "available"
    assert zero["sample_quantity_sum"] == 0
    assert zero["observed_row_count"] == 1
    assert all(item["stock_code"] != "NOT-PRESENT" for item in result["products"])


def test_source_cancellation_marker_is_derived_without_changing_invoice_text(
    tmp_path: Path,
) -> None:
    selected = [2, 11, 12, 13, 14, 15, 16, 143]
    workbook, manifest, database, published = local_paths(tmp_path, selected)
    result = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    observed = {
        item["stock_code"]: item["contributing_rows"][0]
        for item in result["products"]
    }
    assert observed["85123A"]["invoice_no"] == "536365"
    assert observed["85123A"]["invoice_no_source_kind"] == "integer"
    assert observed["85123A"]["source_cancellation_marker"] == "false"
    assert observed["D"]["invoice_no"] == "C536379"
    assert observed["D"]["source_cancellation_marker"] == "true"
    assert observed["LOWER-C"]["invoice_no"] == " c-test"
    assert observed["LOWER-C"]["invoice_no_source_kind"] == "text"
    assert observed["LOWER-C"]["source_cancellation_marker"] == "true"
    assert observed["MISSING"]["invoice_no"] is None
    assert observed["MISSING"]["invoice_no_source_kind"] == "missing"
    assert observed["MISSING"]["source_cancellation_marker"] == "unknown"
    assert observed["BLANK"]["invoice_no"] == "   "
    assert observed["BLANK"]["source_cancellation_marker"] == "unknown"
    assert observed["UNSUPPORTED"]["invoice_no"] == "True"
    assert observed["UNSUPPORTED"]["invoice_no_source_kind"] == "unsupported"
    assert observed["UNSUPPORTED"]["source_cancellation_marker"] == "unknown"
    assert observed["TAB-BLANK"]["invoice_no"] == "\t"
    assert observed["TAB-BLANK"]["source_cancellation_marker"] == "unknown"
    assert observed["TAB-C"]["invoice_no"] == "\tC123\n"
    assert observed["TAB-C"]["source_cancellation_marker"] == "true"
    assert result["transformation_rule_id"].startswith("sha256:")
    assert "archive.ics.uci.edu" in result["source_cancellation_marker_contract"][
        "official_definition_url"
    ]


def test_source_marker_breakdown_reconciles_without_changing_product_grain(
    tmp_path: Path,
) -> None:
    selected = [17, 18, 19, 20]
    workbook, manifest, database, published = local_paths(tmp_path, selected)
    result = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    observed_product = product(result, "MIX")
    assert observed_product["sample_quantity_sum"] == 0
    assert observed_product["observed_row_count"] == 4
    observed = marker_breakdown(observed_product)
    assert observed == {
        "true": {
            "source_cancellation_marker": "true",
            "signed_quantity_sum_text": "-7",
            "observed_row_count": 2,
            "contributing_excel_rows": [18, 20],
        },
        "false": {
            "source_cancellation_marker": "false",
            "signed_quantity_sum_text": "5",
            "observed_row_count": 1,
            "contributing_excel_rows": [17],
        },
        "unknown": {
            "source_cancellation_marker": "unknown",
            "signed_quantity_sum_text": "2",
            "observed_row_count": 1,
            "contributing_excel_rows": [19],
        },
    }
    assert {
        row["description"] for row in observed_product["contributing_rows"]
    } == {
        "MIX DESCRIPTION A",
        "MIX DESCRIPTION B",
        "MIX DESCRIPTION C",
        "MIX DESCRIPTION D",
    }
    assert result["source_cancellation_marker_breakdown_rule_id"].startswith(
        "sha256:"
    )
    assert (
        "Description remains on the contributing source rows"
        in result["applied_source_cancellation_marker_breakdown_sql"]
    )


def test_source_marker_breakdown_keeps_large_sql_sum_as_exact_text(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [21])
    result = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )

    observed_product = product(result, "BIG")
    assert observed_product["sample_quantity_sum"] == 9007199254740993
    assert observed_product["sample_quantity_sum_text"] == "9007199254740993"
    assert observed_product["contributing_rows"][0][
        "quantity_integer_text"
    ] == "9007199254740993"
    assert marker_breakdown(observed_product)["false"][
        "signed_quantity_sum_text"
    ] == "9007199254740993"

    repeated = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    comparison = compare_published_results(
        database,
        published,
        base_run_id=result["run"]["run_id"],
        current_run_id=repeated["run"]["run_id"],
    )
    observed_comparison = comparison["groups"][0]
    assert observed_comparison["base_sample_quantity_sum_text"] == "9007199254740993"
    assert observed_comparison["current_sample_quantity_sum_text"] == "9007199254740993"
    assert observed_comparison["quantity_sum_changed"] is False


def test_actual_uci_marker_breakdown_explains_five_without_calling_it_sales(
    tmp_path: Path,
) -> None:
    source = PROJECT_ROOT / "data" / "source" / "Online Retail.xlsx"
    manifest = PROJECT_ROOT / "config" / "source.json"
    if not source.exists():
        pytest.skip("local preserved UCI workbook is not available")
    source_sha_before = hashlib.sha256(source.read_bytes()).hexdigest()

    result = run_pipeline(
        source_path=source,
        manifest_path=manifest,
        database_path=tmp_path / "actual-marker-breakdown.duckdb",
        published_root=tmp_path / "published",
        selected_rows=[143, 156, 202],
    )

    ducks = product(result, "35004C")
    assert ducks["sample_quantity_sum"] == 5
    assert ducks["observed_row_count"] == 2
    assert marker_breakdown(ducks) == {
        "true": {
            "source_cancellation_marker": "true",
            "signed_quantity_sum_text": "-1",
            "observed_row_count": 1,
            "contributing_excel_rows": [156],
        },
        "false": {
            "source_cancellation_marker": "false",
            "signed_quantity_sum_text": "6",
            "observed_row_count": 1,
            "contributing_excel_rows": [202],
        },
    }
    assert "unknown" not in marker_breakdown(ducks)
    assert [
        row["source_row_number"] for row in ducks["contributing_rows"]
    ] == [156, 202]
    assert [
        row["quantity_integer_text"] for row in ducks["contributing_rows"]
    ] == ["-1", "6"]
    discount = product(result, "D")
    assert discount["sample_quantity_sum"] == -1
    assert marker_breakdown(discount)["true"]["contributing_excel_rows"] == [143]

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook["Online Retail"]
        assert sum(int(sheet.cell(row=row_number, column=4).value) for row_number in [156, 202]) == 5
        assert {
            str(sheet.cell(row=row_number, column=2).value)
            for row_number in [156, 202]
        } == {"35004C"}
        assert str(sheet.cell(row=156, column=1).value).startswith("C")
        assert not str(sheet.cell(row=202, column=1).value).startswith("C")
        assert sheet.cell(row=156, column=3).value == sheet.cell(row=202, column=3).value
    finally:
        workbook.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha_before


def test_actual_uci_sample_matches_independent_workbook_calculation(tmp_path: Path) -> None:
    project_root = Path(__file__).resolve().parents[1]
    source = project_root / "data" / "source" / "Online Retail.xlsx"
    manifest = project_root / "config" / "source.json"
    if not source.exists():
        pytest.skip("local preserved UCI workbook is not available")

    result = run_pipeline(
        source_path=source,
        manifest_path=manifest,
        database_path=tmp_path / "actual.duckdb",
        published_root=tmp_path / "published",
    )

    config = json.loads(manifest.read_text(encoding="utf-8"))
    selected = set(config["selected_excel_rows"])
    independent_totals: dict[str, dict[str, int]] = {}
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[config["sheet"]]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_row=max(selected), max_col=8, values_only=True),
            start=2,
        ):
            if row_number not in selected:
                continue
            key = str(row[1])
            current = independent_totals.setdefault(key, {"quantity": 0, "rows": 0})
            current["quantity"] += int(row[3])
            current["rows"] += 1
    finally:
        workbook.close()

    for item in result["products"]:
        expected = independent_totals[item["stock_code"]]
        assert item["sample_quantity_sum"] == expected["quantity"]
        assert item["observed_row_count"] == expected["rows"]

    assert product(result, "85123A")["sample_quantity_sum"] == 6
    assert product(result, "D")["sample_quantity_sum"] == -1
    assert product(result, "85123A")["contributing_rows"][0][
        "source_cancellation_marker"
    ] == "false"
    assert product(result, "D")["contributing_rows"][0][
        "source_cancellation_marker"
    ] == "true"


def test_run_scope_is_separate_from_preserved_source_rows_and_can_be_compared(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])

    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[3],
    )
    third = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[3, 2],
    )

    assert first["source"]["selected_excel_rows"] == [2, 3]
    assert second["source"]["selected_excel_rows"] == [3]
    assert third["source"]["selected_excel_rows"] == [2, 3]
    assert first["source"]["input_scope_id"] == third["source"]["input_scope_id"]
    assert product(first, "85123A")["sample_quantity_sum"] == 6
    assert [item["stock_code"] for item in second["products"]] == ["71053"]
    assert product(third, "85123A")["sample_quantity_sum"] == 6

    connection = duckdb.connect(str(database), read_only=True)
    try:
        assert connection.execute("SELECT COUNT(*) FROM source_files").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM source_rows").fetchone()[0] == 2
        assert connection.execute("SELECT COUNT(*) FROM run_input_scopes").fetchone()[0] == 3
        assert connection.execute("SELECT COUNT(*) FROM run_input_rows").fetchone()[0] == 5
    finally:
        connection.close()

    comparison = compare_published_results(
        database,
        published,
        base_run_id=first["run"]["run_id"],
        current_run_id=second["run"]["run_id"],
    )
    assert comparison["change_basis"]["code"] == "selection_scope_changed"
    assert comparison["change_basis"]["selected_rows_added"] == []
    assert comparison["change_basis"]["selected_rows_removed"] == [2]
    changed = next(
        item for item in comparison["groups"] if item["stock_code"] == "85123A"
    )
    assert changed["base_sample_quantity_sum"] == 6
    assert changed["current_state"] == "not_present"
    assert changed["removed_source_rows"] == [
        {
            "source_file_id": first["source"]["source_file_id"],
            "sheet": "Online Retail",
            "source_row_number": 2,
        }
    ]
    unchanged = next(
        item for item in comparison["groups"] if item["stock_code"] == "71053"
    )
    assert unchanged["base_sample_quantity_sum"] == 6
    assert unchanged["current_sample_quantity_sum"] == 6
    assert unchanged["changed"] is False

    same_scope_comparison = compare_published_results(
        database,
        published,
        base_run_id=first["run"]["run_id"],
        current_run_id=third["run"]["run_id"],
    )
    assert same_scope_comparison["change_basis"]["code"] == "same_input_scope_and_rule"
    assert all(item["changed"] is False for item in same_scope_comparison["groups"])

    results = list_published_results(database, published)
    assert [item["run_id"] for item in results] == [
        third["run"]["run_id"],
        second["run"]["run_id"],
        first["run"]["run_id"],
    ]
    assert results[0]["is_current"] is True


def test_comparison_keeps_equal_value_and_changed_source_locations_separate(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 10])
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2],
    )
    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[10],
    )

    comparison = compare_published_results(
        database,
        published,
        base_run_id=first["run"]["run_id"],
        current_run_id=second["run"]["run_id"],
    )
    group = comparison["groups"][0]

    assert group["base_sample_quantity_sum"] == 6
    assert group["current_sample_quantity_sum"] == 6
    assert group["base_observed_row_count"] == 1
    assert group["current_observed_row_count"] == 1
    assert [row["source_row_number"] for row in group["removed_source_rows"]] == [2]
    assert [row["source_row_number"] for row in group["added_source_rows"]] == [10]
    assert group["changed"] is True


def test_saved_results_open_their_own_rows_sql_and_exact_comparison_values(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(
        tmp_path, [143, 156, 202, 299]
    )
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[143, 156, 202],
    )
    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[143, 156, 202, 299],
    )
    database_before_reads = hashlib.sha256(database.read_bytes()).hexdigest()

    opened_first = load_published_result(
        database, published, first["run"]["run_id"]
    )
    opened_second = load_published_result(
        database, published, second["run"]["run_id"]
    )

    assert opened_first["is_current"] is False
    assert opened_second["is_current"] is True
    assert opened_first["result"]["run"]["run_id"] == first["run"]["run_id"]
    assert opened_second["result"]["run"]["run_id"] == second["run"]["run_id"]
    first_ducks = product(opened_first["result"], "35004C")
    second_ducks = product(opened_second["result"], "35004C")
    assert first_ducks["sample_quantity_sum_text"] == "5"
    assert [
        row["source_row_number"] for row in first_ducks["contributing_rows"]
    ] == [156, 202]
    assert second_ducks["sample_quantity_sum_text"] == "53"
    assert [
        (row["source_row_number"], row["quantity_integer_text"])
        for row in second_ducks["contributing_rows"]
    ] == [(156, "-1"), (202, "6"), (299, "48")]
    assert opened_first["result"]["applied_sql"] == first["applied_sql"]
    assert opened_second["result"]["applied_sql"] == second["applied_sql"]

    comparison = compare_published_results(
        database,
        published,
        base_run_id=first["run"]["run_id"],
        current_run_id=second["run"]["run_id"],
    )
    ducks_change = next(
        item for item in comparison["groups"] if item["stock_code"] == "35004C"
    )
    assert ducks_change["base_sample_quantity_sum_text"] == "5"
    assert ducks_change["current_sample_quantity_sum_text"] == "53"
    assert ducks_change["quantity_sum_changed"] is True
    assert ducks_change["observed_row_count_changed"] is True
    assert [
        location["source_row_number"]
        for location in ducks_change["added_source_rows"]
    ] == [299]
    assert load_current_result(database, published)["run"]["run_id"] == second["run"]["run_id"]
    assert hashlib.sha256(database.read_bytes()).hexdigest() == database_before_reads


def test_saved_result_read_fails_closed_for_unpublished_states_and_paths(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2])
    current = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    source_file_id = current["source"]["source_file_id"]
    current_artifact = published / f"{current['run']['run_id']}.json"
    orphan_id = "orphan-not-registered"
    (published / f"{orphan_id}.json").write_bytes(current_artifact.read_bytes())

    connection = duckdb.connect(str(database))
    try:
        for run_id, status in (("failed-run", "failed"), ("running-run", "running")):
            connection.execute(
                """
                INSERT INTO pipeline_runs
                    (run_id, source_file_id, input_path, status, started_at)
                VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
                """,
                [run_id, source_file_id, str(workbook), status],
            )
            connection.execute(
                """
                INSERT INTO published_results
                    (run_id, artifact_path, published_at, is_current)
                VALUES (?, ?, CURRENT_TIMESTAMP, FALSE)
                """,
                [run_id, str(current_artifact)],
            )
    finally:
        connection.close()

    for run_id in ("missing-run", orphan_id, "failed-run", "running-run"):
        with pytest.raises(ResultNotFound, match="successful published result not found"):
            load_published_result(database, published, run_id)

    external = tmp_path / "outside.json"
    external.write_bytes(current_artifact.read_bytes())
    connection = duckdb.connect(str(database))
    try:
        connection.execute(
            "UPDATE published_results SET artifact_path = ? WHERE run_id = ?",
            [str(external), current["run"]["run_id"]],
        )
    finally:
        connection.close()
    with pytest.raises(ResultReadError, match="outside the result directory"):
        load_published_result(database, published, current["run"]["run_id"])


@pytest.mark.parametrize("artifact_state", ["missing", "corrupt", "wrong_run_id"])
def test_saved_result_read_does_not_substitute_current_for_bad_artifact(
    tmp_path: Path, artifact_state: str
) -> None:
    case_root = tmp_path / artifact_state
    case_root.mkdir()
    workbook, manifest, database, published = local_paths(case_root, [2, 3])
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2],
    )
    second = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    first_artifact = published / f"{first['run']['run_id']}.json"
    if artifact_state == "missing":
        first_artifact.unlink()
        expected = "published result could not be read"
    elif artifact_state == "corrupt":
        first_artifact.write_text("{not-json", encoding="utf-8")
        expected = "published result could not be read"
    else:
        mismatched = deepcopy(first)
        mismatched["run"]["run_id"] = "different-run"
        first_artifact.write_text(json.dumps(mismatched), encoding="utf-8")
        expected = "does not match the requested run"

    with pytest.raises(ResultReadError, match=expected):
        load_published_result(database, published, first["run"]["run_id"])
    assert load_current_result(database, published)["run"]["run_id"] == second["run"]["run_id"]


def test_failed_new_scope_does_not_replace_or_become_comparable_result(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    first = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )

    with pytest.raises(PipelineFailure, match="simulated failure"):
        run_pipeline(
            source_path=workbook,
            manifest_path=manifest,
            database_path=database,
            published_root=published,
            selected_rows=[3],
            failure_point="after_artifact",
        )

    assert load_current_result(database, published)["run"]["run_id"] == first["run"]["run_id"]
    assert [item["run_id"] for item in list_published_results(database, published)] == [
        first["run"]["run_id"]
    ]


def test_file_or_sql_change_is_not_labeled_as_selection_only(tmp_path: Path) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2])
    base = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
    )
    changed = deepcopy(base)
    changed["run"]["run_id"] = "different-result"
    changed["source"]["source_file_id"] = "sha256:different-source"
    changed["source"]["sha256"] = "different-source"
    changed["source"]["input_scope_id"] = "sha256:different-scope"
    changed["aggregation_rule_id"] = "sha256:different-rule"

    comparison = compare_results(base, changed)

    assert comparison["change_basis"]["code"] == "source_sheet_or_rule_changed"
    assert comparison["change_basis"]["source_file_changed"] is True
    assert comparison["change_basis"]["aggregation_rule_changed"] is True

    rule_only_changed = deepcopy(base)
    rule_only_changed["run"]["run_id"] = "different-rule"
    rule_only_changed["aggregation_rule_id"] = "sha256:different-rule"
    rule_only_comparison = compare_results(base, rule_only_changed)
    assert rule_only_comparison["change_basis"]["aggregation_rule_changed"] is True
    assert rule_only_comparison["groups"][0]["changed"] is False

    value_only_changed = deepcopy(base)
    value_only_changed["run"]["run_id"] = "different-value"
    value_only_changed["products"][0]["sample_quantity_sum"] = 7
    value_only_comparison = compare_results(base, value_only_changed)
    assert value_only_comparison["change_basis"]["code"] == "same_input_scope_and_rule"
    assert value_only_comparison["groups"][0]["changed"] is True
    assert value_only_comparison["groups"][0]["added_source_rows"] == []
    assert value_only_comparison["groups"][0]["removed_source_rows"] == []

    transformation_changed = deepcopy(base)
    transformation_changed["run"]["run_id"] = "different-transformation"
    transformation_changed["transformation_rule_id"] = "sha256:different-transformation"
    transformation_comparison = compare_results(base, transformation_changed)
    assert transformation_comparison["change_basis"]["code"] == "source_sheet_or_rule_changed"
    assert transformation_comparison["change_basis"]["transformation_rule_changed"] is True

    legacy = deepcopy(base)
    legacy["run"]["run_id"] = "legacy-result"
    legacy.pop("transformation_rule_id")
    legacy.pop("source_cancellation_marker_contract")
    legacy.pop("applied_transformation_sql")
    for item in legacy["products"]:
        for row in item["contributing_rows"]:
            row.pop("source_cancellation_marker")
    legacy_comparison = compare_results(legacy, base)
    assert legacy_comparison["change_basis"]["transformation_rule_changed"] is True


@pytest.mark.parametrize("checkpoint", PROCESS_CHECKPOINTS)
def test_forced_process_termination_recovers_from_database_evidence(
    tmp_path: Path, checkpoint: str
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    baseline = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2],
    )
    process, barrier = start_crash_worker(
        workbook=workbook,
        manifest=manifest,
        database=database,
        published=published,
        selected_rows=[3],
        checkpoint=checkpoint,
    )
    try:
        kill_owned_worker(process)
    finally:
        if process.poll() is None:
            process.kill()
            process.wait(timeout=10)

    crashed_run_id = barrier["run_id"]
    current_after_death = load_current_result(database, published)
    expected_current = (
        crashed_run_id
        if checkpoint == "after_publish_commit"
        else baseline["run"]["run_id"]
    )
    assert current_after_death["run"]["run_id"] == expected_current

    before_recovery = {
        item["run_id"]: item for item in load_run_history(database)
    }[crashed_run_id]
    expected_status = (
        "succeeded" if checkpoint == "after_publish_commit" else "running"
    )
    assert before_recovery["status"] == expected_status
    assert before_recovery["recovered_at"] is None

    connection = duckdb.connect(str(database), read_only=True)
    try:
        rows_after_death = connection.execute(
            "SELECT COUNT(*) FROM source_rows"
        ).fetchone()[0]
        crashed_membership = connection.execute(
            "SELECT COUNT(*) FROM run_input_rows WHERE run_id = ?",
            [crashed_run_id],
        ).fetchone()[0]
    finally:
        connection.close()
    if checkpoint == "before_data_commit":
        assert rows_after_death == 1
        assert crashed_membership == 0
    else:
        assert rows_after_death == 2
        assert crashed_membership == 1

    active_artifact = published / f"{crashed_run_id}.json"
    assert active_artifact.exists() is (
        checkpoint in {"after_artifact_write", "after_publish_commit"}
    )
    successful_before_recovery = {
        item["run_id"] for item in list_published_results(database, published)
    }
    assert baseline["run"]["run_id"] in successful_before_recovery
    assert (crashed_run_id in successful_before_recovery) is (
        checkpoint == "after_publish_commit"
    )

    retried = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    assert {item["stock_code"] for item in retried["products"]} == {
        "71053",
        "85123A",
    }
    assert product(retried, "85123A")["sample_quantity_sum"] == 6
    assert product(retried, "71053")["sample_quantity_sum"] == 6
    assert load_current_result(database, published)["run"]["run_id"] == retried[
        "run"
    ]["run_id"]

    history = {item["run_id"]: item for item in load_run_history(database)}
    recovered = history[crashed_run_id]
    successful_after_recovery = {
        item["run_id"] for item in list_published_results(database, published)
    }
    if checkpoint == "after_publish_commit":
        assert recovered["status"] == "succeeded"
        assert recovered["recovered_at"] is None
        assert crashed_run_id in successful_after_recovery
        assert active_artifact.is_file()
    else:
        assert recovered["status"] == "failed"
        assert recovered["recovered_at"] is not None
        assert "previous writer process ended" in recovered["recovery_reason"]
        assert crashed_run_id not in successful_after_recovery
        if checkpoint == "after_artifact_write":
            orphaned = published / "orphaned" / f"{crashed_run_id}.json"
            assert not active_artifact.exists()
            assert orphaned.is_file()
            assert recovered["recovered_artifact_path"] == str(orphaned.resolve())
        else:
            assert recovered["recovered_artifact_path"] is None

    connection = duckdb.connect(str(database), read_only=True)
    try:
        assert connection.execute("SELECT COUNT(*) FROM source_files").fetchone()[0] == 1
        assert connection.execute("SELECT COUNT(*) FROM source_rows").fetchone()[0] == 2
        assert connection.execute(
            "SELECT COUNT(*) FROM published_results WHERE is_current = TRUE"
        ).fetchone()[0] == 1
    finally:
        connection.close()


def test_live_writer_lock_rejects_second_writer_without_recovering_it(
    tmp_path: Path,
) -> None:
    workbook, manifest, database, published = local_paths(tmp_path, [2, 3])
    run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2],
    )
    process, barrier = start_crash_worker(
        workbook=workbook,
        manifest=manifest,
        database=database,
        published=published,
        selected_rows=[3],
        checkpoint="after_data_commit",
    )
    try:
        with pytest.raises(PipelineBusy, match="another local pipeline writer"):
            run_pipeline(
                source_path=workbook,
                manifest_path=manifest,
                database_path=database,
                published_root=published,
                selected_rows=[2, 3],
            )
        lock_state = json.loads(
            database.with_name(f"{database.name}.writer.lock").read_text(
                encoding="utf-8"
            )
        )
        assert lock_state["owner_pid"] == process.pid
        assert process.poll() is None
    finally:
        kill_owned_worker(process)

    child_history = {
        item["run_id"]: item for item in load_run_history(database)
    }[barrier["run_id"]]
    assert child_history["status"] == "running"
    assert child_history["recovered_at"] is None

    retried = run_pipeline(
        source_path=workbook,
        manifest_path=manifest,
        database_path=database,
        published_root=published,
        selected_rows=[2, 3],
    )
    recovered = {
        item["run_id"]: item for item in load_run_history(database)
    }[barrier["run_id"]]
    assert recovered["status"] == "failed"
    assert recovered["recovered_at"] is not None
    assert product(retried, "85123A")["sample_quantity_sum"] == 6
    assert product(retried, "71053")["sample_quantity_sum"] == 6
